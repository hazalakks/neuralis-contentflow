"""
Builds the submission deck as an editable PowerPoint file that imports cleanly into Canva.

  python scripts/build_presentation.py
  -> presentation/NEURALIS-CONTENTFLOW.pptx

Canva: Create a design -> Import file (or drag the .pptx onto the Canva home page).
Fonts used (both available in Canva): Space Grotesk (headlines), Inter (body).
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "assets" / "images"
OUT = ROOT / "presentation" / "NEURALIS-CONTENTFLOW.pptx"

HEAD, BODY, MONO = "Space Grotesk", "Inter", "Consolas"


def rgb(h): return RGBColor.from_string(h)


BG, SURF, SURF2 = rgb("0B0F14"), rgb("121820"), rgb("1A222D")
CYAN, EMER, VIOL = rgb("00F2FE"), rgb("10B981"), rgb("8B5CF6")
WHITE, MUTED, DIM = rgb("F8FAFC"), rgb("94A3B8"), rgb("64748B")
AMBER, RED = rgb("F59E0B"), rgb("F87171")
LINE = rgb("263040")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 12


# ----------------------------------------------------------------- helpers
def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def rect(s, x, y, w, h, fill=SURF, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def tb(s, x, y, w, h, paras, size=14, color=WHITE, font=BODY, bold=False,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4, italic=False):
    """paras: str | list of (str | list of runs). run: str | (text, dict(size,color,bold,font,italic))"""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space)
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            text, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = text
            f = run.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return box


def image_fit(s, path, x, y, w, h):
    """Insert an image cropped to fill the box (object-fit: cover)."""
    iw, ih = Image.open(path).size
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        c = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def header(s, n, kicker, title):
    tb(s, 0.6, 0.45, 9, 0.3, kicker.upper(), size=11, color=CYAN, font=HEAD, bold=True)
    tb(s, 0.6, 0.75, 12, 0.8, title, size=30, font=HEAD, bold=True)
    tb(s, 0.6, 7.05, 6, 0.25, "NEURALIS · CONTENTFLOW", size=9, color=DIM, font=HEAD, bold=True)
    tb(s, 11.73, 7.05, 1, 0.25, f"{n:02d} / {TOTAL}", size=9, color=DIM, font=HEAD, align=PP_ALIGN.RIGHT)


def label(s, x, y, text, color=CYAN, w=5):
    tb(s, x, y, w, 0.25, text.upper(), size=10, color=color, font=HEAD, bold=True)


def chip(s, x, y, text, color=CYAN, w=None, size=11):
    w = w or (0.2 + 0.085 * len(text) * size / 11)
    rect(s, x, y, w, 0.34, fill=SURF2, line=LINE, radius=0.5)
    tb(s, x, y, w, 0.34, text, size=size, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def set_borders(cell, color="263040", width=6350):
    from pptx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    for idx, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(width), "cmpd": "sng"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
        tcPr.insert(idx, ln)  # schema order: lnL, lnR, lnT, lnB, then fill


def table(s, x, y, col_w, rows, header_color=CYAN, size=11, row_h=0.42, first_bold=True, colors=None):
    shape = s.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y),
                               Inches(sum(col_w)), Inches(row_h * len(rows)))
    t = shape.table
    for i, w in enumerate(col_w):
        t.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURF2 if r == 0 else (SURF if r % 2 else BG)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            set_borders(cell)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = val
            f = run.font
            f.name = HEAD if r == 0 else BODY
            f.size = Pt(size - 1 if r == 0 else size)
            f.bold = r == 0 or (first_bold and c == 0)
            col = header_color if r == 0 else WHITE if c == 0 else rgb("CBD5E1")
            if colors and r > 0 and (r, c) in colors:
                col = colors[(r, c)]
            f.color.rgb = col
    return shape


# ----------------------------------------------------------------- 1 · cover
s = new_slide()
image_fit(s, IMG / "hero-workspace.jpg", 0, 0, 13.333, 7.5)
ov = rect(s, 0, 0, 13.333, 7.5, fill=BG, shape=MSO_SHAPE.RECTANGLE)
ov.fill.fore_color.rgb = BG
from pptx.oxml.ns import qn  # noqa: E402  (transparency for overlay)
sf = ov.fill._xPr.find(qn("a:solidFill"))
clr = sf[0]
alpha = clr.makeelement(qn("a:alpha"), {"val": "35000"})
clr.append(alpha)
rect(s, 0, 0, 7.4, 7.5, fill=BG, shape=MSO_SHAPE.RECTANGLE)
tb(s, 0.8, 1.2, 6, 0.3, "FINAL PROJECT · AI CONTENT PRODUCTION · 2026", size=11, color=CYAN, font=HEAD, bold=True)
tb(s, 0.8, 1.8, 6.4, 1.0, "NEURALIS", size=60, font=HEAD, bold=True, color=CYAN)
tb(s, 0.8, 2.85, 6.4, 0.7, "CONTENTFLOW", size=34, font=HEAD, bold=True)
tb(s, 0.8, 3.75, 6.0, 1.2,
   "An end-to-end, AI-powered content system for a focus-tech brand: from strategy to a full week of blog, social, image, video and voice content.",
   size=16, color=rgb("CBD5E1"), space=0)
tb(s, 0.8, 5.6, 6, 0.35, "[Your Name]", size=16, font=HEAD, bold=True)
tb(s, 0.8, 6.0, 6, 0.3, "Yapay Zeka ile İçerik Üretimi Eğitimi — Final Projesi", size=12, color=MUTED)
tb(s, 0.8, 6.85, 6, 0.3, "NEURALIS is a fictional concept brand.", size=9, color=DIM)

# ----------------------------------------------------------------- 2 · brief & brand
s = new_slide()
header(s, 2, "Stage 01 · Brand analysis", "The brief & the brand")
rect(s, 0.6, 1.75, 5.9, 3.3)
label(s, 0.95, 2.0, "The brief")
tb(s, 0.95, 2.4, 5.2, 2.5,
   ["A software company wants to reduce its marketing team's weekly content workload and keep its brand voice consistent across every channel.",
    [("Goal: ", {"bold": True, "color": WHITE}), ("design an end-to-end content system with AI tools: strategy → calendar → text, image, video and voice → review.", {"color": rgb("CBD5E1")})]],
   size=15, space=12)
rect(s, 6.83, 1.75, 5.9, 3.3)
label(s, 7.18, 2.0, "The brand")
tb(s, 7.18, 2.4, 5.3, 0.5, "NEURALIS", size=26, font=HEAD, bold=True, color=CYAN)
tb(s, 7.18, 3.0, 5.3, 2.0,
   ["Neuro-adaptive focus audio for knowledge workers.",
    [("Sector  ", {"color": MUTED}), "Cognitive wellness tech · Productivity SaaS"],
    [("Model  ", {"color": MUTED}), "B2C subscription · B2B team plans"],
    [("Promise  ", {"color": MUTED}), ("Focus, by design.", {"bold": True})]],
   size=14, space=8)
rect(s, 0.6, 5.3, 12.13, 1.45, fill=SURF2, line=CYAN)
label(s, 0.95, 5.5, "Positioning statement")
tb(s, 0.95, 5.85, 11.5, 0.8,
   "For knowledge workers who lose their best hours to constant switching, NEURALIS is the focus-audio platform that adapts to how you work, not just what time it is, so deep work becomes the default, not a fight.",
   size=15, italic=True)

# ----------------------------------------------------------------- 3 · competitors
s = new_slide()
header(s, 3, "Stage 01 · Competitor analysis", "Four competitors, one clear gap")
comps = [
    ("Brain.fm", "Functional music. Publishes peer-reviewed research.", "Sessions are chosen manually. Sound doesn't react to what you're doing."),
    ("Endel", "Adapts to time of day, weather and heart rate. Strong design.", "Broad wellbeing focus. No work-context signals, no team layer."),
    ("Focus@Will", "Curated focus-music channels by listener type.", "Static playlists, little personalisation."),
    ("Freedom", "Cross-device app & site blocking. Clear value.", "Only removes distractions. Doesn't improve the working state."),
]
for i, (n, good, gap) in enumerate(comps):
    x = 0.6 + i * 3.08
    rect(s, x, 1.75, 2.9, 3.85)
    tb(s, x + 0.3, 2.0, 2.4, 0.5, n, size=20, font=HEAD, bold=True)
    label(s, x + 0.3, 2.65, "Does well", EMER)
    tb(s, x + 0.3, 2.95, 2.35, 1.1, good, size=12, color=rgb("CBD5E1"))
    label(s, x + 0.3, 4.05, "Gap", AMBER)
    tb(s, x + 0.3, 4.35, 2.35, 1.2, gap, size=12, color=rgb("CBD5E1"))
rect(s, 0.6, 5.85, 12.13, 0.95, fill=SURF2, line=CYAN)
tb(s, 0.95, 5.85, 11.5, 0.95,
   [[("NEURALIS opportunity → ", {"bold": True, "color": CYAN}),
     "audio that adapts to work-context signals (typing cadence, session phase, calendar load), plus a team layer for B2B."]],
   size=15, anchor=MSO_ANCHOR.MIDDLE)

# ----------------------------------------------------------------- 4 · identity
s = new_slide()
header(s, 4, "Stage 01 · Brand identity card", "Identity: calm, precise, evidence-first")
label(s, 0.6, 1.75, "Values")
x = 0.6
for v, w in [("Evidence over hype", 1.55), ("Calm over urgency", 1.5), ("Sustainable performance", 2.0), ("Respect for attention", 1.85)]:
    x += chip(s, x, 2.05, v, WHITE, w=w, size=10) + 0.1
label(s, 0.6, 2.75, "Palette")
pal = [("0B0F14", "Obsidian"), ("00F2FE", "Synapse Cyan"), ("10B981", "Bio-Emerald"), ("8B5CF6", "Pulse Violet"), ("F8FAFC", "Neural White")]
for i, (hx, nm) in enumerate(pal):
    xx = 0.6 + i * 1.5
    rect(s, xx, 3.05, 1.35, 0.9, fill=rgb(hx), line=LINE, radius=0.12)
    tb(s, xx, 4.02, 1.4, 0.25, f"#{hx}", size=11, font=HEAD, bold=True)
    tb(s, xx, 4.27, 1.4, 0.25, nm, size=10, color=MUTED)
label(s, 0.6, 4.8, "Typography")
tb(s, 0.6, 5.1, 3.6, 0.5, "Space Grotesk", size=24, font=HEAD, bold=True)
tb(s, 0.6, 5.6, 3.6, 0.3, "Headlines · technical, precise", size=11, color=MUTED)
tb(s, 4.0, 5.1, 3.4, 0.5, "Inter", size=24, font=BODY)
tb(s, 4.0, 5.6, 3.4, 0.3, "Body · high legibility", size=11, color=MUTED)
tb(s, 0.6, 6.1, 7.2, 0.8,
   [[("Visual style  ", {"bold": True, "color": EMER}), "dark cyber-minimal UI, one soundwave motif, soft volumetric light"],
    [("Avoid  ", {"bold": True, "color": RED}), "clinical imagery, robot clichés, neon clutter, real brand logos"]],
   size=11, space=4)
image_fit(s, IMG / "brainwave-synapses.jpg", 8.2, 1.75, 4.53, 5.0)

# ----------------------------------------------------------------- 5 · why who what how + personas
s = new_slide()
header(s, 5, "Stage 02 · Audience & strategy", "Why · Who · What · How")
quad = [("WHY", "Attention is the scarcest resource in modern work. We exist to protect it.", CYAN),
        ("WHO", "Engineers, analysts, designers and founders doing 4+ hours of demanding screen work a day.", EMER),
        ("WHAT", "Adaptive focus audio + session insights that help people enter and sustain deep work.", VIOL),
        ("HOW", "Evidence-based, calm and precise communication. Never hype, never hustle.", AMBER)]
for i, (k, v, c) in enumerate(quad):
    x, y = 0.6 + (i % 2) * 3.65, 1.75 + (i // 2) * 2.55
    rect(s, x, y, 3.5, 2.4)
    tb(s, x + 0.3, y + 0.25, 3, 0.5, k, size=22, font=HEAD, bold=True, color=c)
    tb(s, x + 0.3, y + 0.85, 2.95, 1.5, v, size=13, color=rgb("CBD5E1"))
personas = [("Deniz, 31", "Senior Software Engineer · primary",
             "Pain: pings and meetings fragment the day; foggy by 4 pm.", "Goal: two protected 90-min deep-work blocks a day.", "LinkedIn · X · Hacker News"),
            ("Selin, 36", "Startup Co-founder · B2B buyer",
             "Pain: decision fatigue; a team that's always on and always tired.", "Goal: a sustainable focus culture for the whole team.", "LinkedIn · newsletters · podcasts")]
for i, (n, role, pain, goal, ch) in enumerate(personas):
    y = 1.75 + i * 2.55
    rect(s, 8.0, y, 4.73, 2.4, fill=SURF2, line=LINE)
    tb(s, 8.3, y + 0.2, 4.2, 0.4, n, size=18, font=HEAD, bold=True)
    tb(s, 8.3, y + 0.6, 4.2, 0.3, role, size=11, color=CYAN)
    tb(s, 8.3, y + 0.95, 4.2, 1.4, [pain, goal, [(ch, {"color": MUTED, "size": 10})]], size=12, color=rgb("CBD5E1"), space=5)

# ----------------------------------------------------------------- 6 · voice guide
s = new_slide()
header(s, 6, "Stage 02 · One-page brand voice guide", "Voice: \"a calm neuroscientist who also ships software\"")
table(s, 0.6, 1.75, [1.7, 3.0, 2.9], [
    ["Trait", "We sound like…", "We never sound like…"],
    ["Evidence-based", "\"A 2005 field study found interrupted work took ~25 min to resume.\"", "\"Scientists PROVE this frequency makes you 10x smarter.\""],
    ["Precise", "Short sentences. Numbers only when sourced.", "Ornamental paragraphs, vague buzzwords."],
    ["Calm", "Reassuring, low-pressure CTAs.", "FOMO, countdowns, \"Don't miss out!!!\""],
    ["Humane", "Rest is part of performance.", "\"Sleep is for the weak.\" \"Grind harder.\""],
], size=11, row_h=0.82)
rect(s, 8.45, 1.75, 4.28, 2.0)
label(s, 8.75, 1.95, "Use", EMER)
tb(s, 8.75, 2.25, 3.8, 0.6, "deep work · focus · research suggests · designed to support · recovery", size=12, color=rgb("CBD5E1"))
label(s, 8.75, 2.85, "Avoid", RED)
tb(s, 8.75, 3.15, 3.8, 0.6, "miracle · hack your brain · 10x · cures · grind · hustle", size=12, color=rgb("CBD5E1"))
rect(s, 8.45, 3.95, 4.28, 2.8, fill=SURF2, line=CYAN)
label(s, 8.75, 4.15, "5 hard rules")
tb(s, 8.75, 4.5, 3.8, 2.2,
   ["1  No medical claims", "2  No unsourced statistics", "3  Hedge emerging science", "4  No artist / voice imitation", "5  No real brand marks in visuals"],
   size=13, space=6)

# ----------------------------------------------------------------- 7 · pillars & calendar
s = new_slide()
header(s, 7, "Stage 03 · Content pillars & 7-day calendar", "One week, one arc: trust → value → product → habit")
pillars = [("Science of Focus", "35%", CYAN), ("Deep Work Systems", "30%", EMER), ("Inside NEURALIS", "20%", VIOL), ("Sustainable Performance", "15%", AMBER)]
for i, (n, p, c) in enumerate(pillars):
    x = 0.6 + i * 3.08
    rect(s, x, 1.7, 2.9, 0.75)
    tb(s, x + 0.2, 1.7, 1.9, 0.75, f"{i+1} · {n}", size=12, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 2.0, 1.7, 0.75, 0.75, p, size=18, font=HEAD, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
table(s, 0.6, 2.65, [1.0, 3.1, 4.6, 3.43], [
    ["Day", "Channel · Format", "Topic", "Why today"],
    ["Mon", "LinkedIn · RTF post", "The real cost of a \"quick question\"", "Planning day: open to systems"],
    ["Tue", "Instagram · 5-slide carousel", "Your brain at work: beta, alpha, theta", "Saveable explainer builds authority"],
    ["Wed", "Blog · SEO article (hub)", "Cognitive fatigue isn't a willpower problem", "Evergreen hub mid-week"],
    ["Thu", "X · 6-post thread", "Why your focus playlist stops working", "Tech audience rewards depth"],
    ["Fri", "Reels / Shorts · 18 s video", "Your focus isn't broken. It's overloaded.", "Meets the Friday energy dip"],
    ["Sat", "Audio · TTS voiceover", "The One-Minute Neural Reset", "Weekend recovery"],
    ["Sun", "Email · weekly digest", "Three rules for a clearer Monday", "Closes the weekly loop"],
], size=11, row_h=0.47, colors={(3, 0): CYAN})
tb(s, 0.6, 6.72, 12, 0.3, "Full calendar with personas, rationale and KPI targets: content-calendar.xlsx", size=10, color=MUTED)

# ----------------------------------------------------------------- 8 · blog
s = new_slide()
header(s, 8, "Stage 03 · Blog post · CO-STAR prompt", "SEO blog: honest science, practical habits")
image_fit(s, IMG / "hero-workspace.jpg", 0.6, 1.75, 5.2, 3.0)
rect(s, 0.6, 4.95, 5.2, 1.8)
label(s, 0.9, 5.12, "SEO setup")
tb(s, 0.9, 5.45, 4.7, 1.3,
   [[("Primary keyword  ", {"color": MUTED}), ("cognitive fatigue", {"bold": True, "color": CYAN})],
    [("Secondary  ", {"color": MUTED}), "deep work · context switching · focus music · binaural beats"],
    [("Length  ", {"color": MUTED}), "~1,050 words · 6 H2s · 1 table · 4 sources"]],
   size=11, space=5)
tb(s, 6.2, 1.75, 6.5, 1.1, "Cognitive Fatigue Isn't a Willpower Problem: How Sound Can Help You Get Back to Deep Work",
   size=20, font=HEAD, bold=True, space=0)
tb(s, 6.2, 2.95, 6.5, 0.8,
   "Meta: Cognitive fatigue isn't a lack of discipline. It's what constant switching does to your attention. Here's what research says, plus three habits to start today.",
   size=11, color=MUTED, italic=True)
tb(s, 6.2, 3.85, 6.5, 1.7,
   [[("→ ", {"color": CYAN}), "Interrupted work took ~25 min on average to resume (Mark et al., 2005)"],
    [("→ ", {"color": CYAN}), "Every task switch has a measurable cost (Rubinstein et al., 2001)"],
    [("→ ", {"color": CYAN}), "Focus-sound evidence is promising, not settled (meta-analysis, 2019)"],
    [("→ ", {"color": CYAN}), "3 habits: batch interruptions · one 90-min block · steady soundscape"]],
   size=13, space=7)
rect(s, 6.2, 5.75, 6.53, 1.0, fill=SURF2, line=LINE)
tb(s, 6.45, 5.85, 6.1, 0.85,
   ["Sources: Mark, González & Harris (2005) · Rubinstein, Meyer & Evans (2001) · Garcia-Argibay et al. (2019) · Woods et al. (2024)",
    [("Includes a \"not a medical device\" disclaimer.", {"color": MUTED})]],
   size=10, space=3)

# ----------------------------------------------------------------- 9 · social
s = new_slide()
header(s, 9, "Stage 03 · 3 social posts · RTF formula", "Social: one idea per platform, written with RTF")
posts = [
    ("LinkedIn", "Mon · Deep Work",
     ["R  Evidence-driven thought leader", "T  Hidden cost of context switching; only the verified 25-min fact", "F  Hook · 3 paragraphs · question · 3 hashtags"],
     "A \"quick question\" is never quick for the person answering it.\n\nWhen work was interrupted, it took around 25 minutes on average to get back to it. (Mark et al., CHI 2005)\n\nWhat's one rule your team uses to protect deep work?"),
    ("Instagram carousel", "Tue · Science",
     ["R  NEURALIS science editor", "T  Explain beta / alpha / theta as associations, not switches", "F  5 slides · hook ≤ 8 words · caption + save CTA"],
     "1  Your brain at work, in 3 states.\n2  BETA ~13–30 Hz: alert, effortful\n3  ALPHA ~8–12 Hz: relaxed wakefulness\n4  THETA ~4–7 Hz: drifting, associative\n5  Focus = fewer interruptions, so your brain can settle."),
    ("X / Twitter thread", "Thu · Product",
     ["R  Senior engineer & community lead", "T  Why static playlists fade; cite 2019 meta-analysis; hedge", "F  6 posts · ≤ 280 chars · no hashtags"],
     "1/6 Your focus playlist works great for 20 minutes. Then you start skipping tracks…\n\n4/6 A 2019 meta-analysis of 22 studies found a moderate effect, but results varied a lot…\n\n5/6 So NEURALIS adapts instead of promising a magic frequency."),
]
for i, (plat, meta, rtf, copy) in enumerate(posts):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.93, 5.0)
    tb(s, x + 0.25, 1.95, 3.4, 0.4, plat, size=17, font=HEAD, bold=True)
    tb(s, x + 0.25, 2.35, 3.4, 0.3, meta, size=10, color=CYAN)
    rect(s, x + 0.2, 2.75, 3.53, 1.35, fill=BG, line=LINE, radius=0.08)
    tb(s, x + 0.35, 2.85, 3.25, 1.2, rtf, size=9.5, font=MONO, color=EMER, space=3)
    tb(s, x + 0.25, 4.3, 3.45, 2.4, copy.split("\n"), size=11, color=rgb("CBD5E1"), space=2)

# ----------------------------------------------------------------- 10 · media
s = new_slide()
header(s, 10, "Stage 03 · Images, video & voiceover", "Images, an 18 s video and a TTS voiceover")
poster = ROOT / "assets" / "video" / "neuralis-teaser-poster.jpg"
video = ROOT / "assets" / "video" / "neuralis-teaser-9x16.mp4"
s.shapes.add_movie(str(video), Inches(0.6), Inches(1.7), Inches(2.98), Inches(5.3),
                   poster_frame_image=str(poster), mime_type="video/mp4")
for i, (f, cap) in enumerate([("hero-workspace.jpg", "Hero · 16:9"), ("brainwave-synapses.jpg", "Brainwave · 1:1"), ("product-mobile-concept.jpg", "Product · 1:1")]):
    x = 3.95 + i * 2.95
    image_fit(s, IMG / f, x, 1.7, 2.8, 1.6)
    tb(s, x, 3.35, 2.8, 0.25, cap, size=10, color=MUTED)
table(s, 3.95, 3.75, [1.0, 3.6, 4.18], [
    ["Time", "On-screen text", "Voiceover"],
    ["0–5 s", "YOUR FOCUS ISN'T BROKEN.", "\"…It's overloaded.\""],
    ["5–10 s", "SILENCE THE NOISE.", "\"Give your mind one thing to hold.\""],
    ["10–15 s", "SOUND THAT ADAPTS TO YOU.", "\"…shapes sound around the way you work.\""],
    ["15–18 s", "NEURALIS · Focus, by design.", "\"Neuralis. Focus, by design.\""],
], size=10, row_h=0.36)
rect(s, 3.95, 5.7, 8.78, 1.3, fill=SURF2, line=CYAN)
tb(s, 4.25, 5.82, 8.3, 1.1,
   [[("🔊  The One-Minute Neural Reset", {"bold": True, "size": 14, "font": HEAD}), ("   49.6 s · Piper neural TTS (en_US-ryan-high) · ambient bed", {"color": MUTED, "size": 10})],
    [("\"You are not failing at focus. You are carrying too much at once. So for the next minute… carry less.\"", {"italic": True, "color": rgb("CBD5E1")})],
    [("[ add neural-reset-voiceover.mp3 here in Canva ]", {"color": AMBER, "size": 9})]],
   size=12, space=4)

# ----------------------------------------------------------------- 11 · review
s = new_slide()
header(s, 11, "Stage 04 · Copyright & hallucination check", "Review: what the AI got wrong, and how it was fixed")
table(s, 0.6, 1.75, [3.0, 2.95, 2.55], [
    ["First draft said", "Problem", "Fix"],
    ["\"23 min 15 s to refocus\"", "Interview figure, not the paper", "Peer-reviewed ~25 min"],
    ["\"2.1 hours lost per day\"", "No source exists", "Removed"],
    ["\"40% more fatigue & errors\"", "Distorted estimate", "Removed"],
    ["\"Alpha state in 15 minutes\"", "Unsupported claim", "Hedged wording"],
    ["Hinted at ADHD / anxiety help", "Medical claim", "Removed + disclaimer"],
    ["Hero image v1", "Real brand logo on monitor", "Regenerated, no logos"],
    ["Voice dataset CC BY-NC-SA", "Non-commercial licence", "Documented"],
], size=11, row_h=0.6, colors={(r, 1): (RED if r in (2, 3, 4, 5) else AMBER) for r in range(1, 8)})
for i, (big, small, c) in enumerate([("20", "issues found & fixed", AMBER), ("4 / 4", "draft statistics failed review", RED), ("0", "unsourced numbers in final content", EMER)]):
    y = 1.75 + i * 1.32
    rect(s, 9.35, y, 3.38, 1.18)
    tb(s, 9.6, y + 0.1, 3, 0.6, big, size=30, font=HEAD, bold=True, color=c)
    tb(s, 9.6, y + 0.72, 3, 0.35, small, size=11, color=MUTED)
rect(s, 9.35, 5.75, 3.38, 1.0, fill=SURF2, line=CYAN)
tb(s, 9.55, 5.82, 3.0, 0.9, "4 checks per piece: tone · keyword · copyright · sources", size=11, anchor=MSO_ANCHOR.MIDDLE)

# ----------------------------------------------------------------- 12 · process
s = new_slide()
header(s, 12, "Process transparency · AI usage report", "Toolchain, the key lesson & what's next")
tools = [("Text", "Gemini (draft) → Claude (revision & fact-check) via Antigravity", CYAN),
         ("Images", "Antigravity built-in image generation (Google model)", EMER),
         ("Voice", "Piper neural TTS, offline & open source", VIOL),
         ("Video", "FFmpeg, fully scripted in Python", AMBER),
         ("Calendar", "Python + openpyxl → .xlsx / .csv", CYAN),
         ("Fact-check", "Primary sources (DOIs) via web search", EMER)]
for i, (k, v, c) in enumerate(tools):
    x, y = 0.6 + (i % 3) * 2.75, 1.75 + (i // 3) * 1.35
    rect(s, x, y, 2.6, 1.2)
    tb(s, x + 0.2, y + 0.15, 2.3, 0.3, k, size=13, font=HEAD, bold=True, color=c)
    tb(s, x + 0.2, y + 0.5, 2.25, 0.7, v, size=10, color=rgb("CBD5E1"))
rect(s, 8.95, 1.75, 3.78, 2.55, fill=SURF2, line=CYAN)
label(s, 9.2, 1.95, "Most effective prompt change")
tb(s, 9.2, 2.3, 3.35, 1.95,
   ["Adding a source whitelist:", [("\"Verified sources you MAY use: … Do NOT introduce other statistics.\"", {"font": MONO, "size": 10, "color": EMER})],
    "This stopped the model from inventing numbers."],
   size=12, space=6)
label(s, 0.6, 4.6, "Automation design (optional practice)")
steps = ["Calendar\n(Sheets)", "AI draft\n+ images", "Auto checks", "Slack\napproval", "Buffer /\nMeta API"]
for i, st in enumerate(steps):
    x = 0.6 + i * 2.45
    sh = rect(s, x, 4.95, 2.3, 0.95, fill=SURF, line=LINE, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    tb(s, x + 0.35, 4.95, 1.65, 0.95, st.split("\n"), size=11, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)
rect(s, 0.6, 6.15, 12.13, 0.65, fill=SURF2, line=LINE)
tb(s, 0.9, 6.15, 11.6, 0.65,
   [[("GitHub  ", {"bold": True, "color": CYAN}), ("github.com/hazalakks/neuralis-contentflow", {"font": MONO}),
     ("      Full docs, prompt library, review log & reproducible media scripts", {"color": MUTED, "size": 11})]],
   size=13, anchor=MSO_ANCHOR.MIDDLE)

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print("Wrote", OUT.relative_to(ROOT))
