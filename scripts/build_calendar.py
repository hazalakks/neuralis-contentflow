"""
Builds the 7-day content calendar as Excel (.xlsx) and CSV from a single source of truth.

  python scripts/build_calendar.py
  -> docs/content-calendar.xlsx
  -> docs/content-calendar.csv
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

HEADERS = [
    "Day", "Date", "Channel", "Format", "Pillar", "Topic / Hook",
    "Target persona", "Rationale (why this, why today)", "Goal / KPI", "Asset",
]

ROWS = [
    ["1 · Mon", "2026-10-19", "LinkedIn", "Text post (RTF)", "2 · Deep Work Systems",
     "The real cost of a 'quick question': why context switching drains your best hours",
     "Deniz (engineer), Selin (founder)",
     "Mondays are planning days: people set up their week and are receptive to systems. Opens the week with a relatable pain point.",
     "Engagement rate ≥ 4% · saves", "content/social-posts.md#post-1--linkedin"],
    ["2 · Tue", "2026-10-20", "Instagram", "5-slide carousel (RTF)", "1 · Science of Focus",
     "Your brain at work: a simple map of beta, alpha and theta states",
     "Deniz",
     "Visual explainer turns abstract neuroscience into a saveable reference. Builds authority before the flagship blog.",
     "Saves & shares", "assets/images/brainwave-synapses.jpg"],
    ["3 · Wed", "2026-10-21", "Blog (website)", "SEO article (CO-STAR)", "1 + 2",
     "Cognitive fatigue isn't a willpower problem: how sound can help you get back to deep work",
     "Deniz, Selin",
     "Flagship evergreen piece mid-week, when search intent for productivity topics is high. Every other post this week links back to it.",
     "Organic sessions · avg. time on page ≥ 3 min · newsletter sign-ups",
     "content/blog-post.md"],
    ["4 · Thu", "2026-10-22", "X / Twitter", "6-post thread (RTF)", "3 · Inside NEURALIS",
     "Why your focus playlist stops working after 30 minutes, and what adaptive audio does differently",
     "Deniz (tech community)",
     "Tech audiences on X reward technical depth. Differentiates the product from static playlists without attacking competitors.",
     "Bookmarks · profile clicks", "content/social-posts.md#post-3--x--twitter-thread"],
    ["5 · Fri", "2026-10-23", "Instagram Reels / YouTube Shorts / TikTok", "18 s vertical video", "4 · Sustainable Performance",
     "Your focus isn't broken. It's overloaded.",
     "Deniz, Selin",
     "Friday afternoon is the week's lowest-energy point. A short, calming video meets people exactly there and drives app trials.",
     "3-sec hold rate · views · link clicks", "assets/video/neuralis-teaser-9x16.mp4"],
    ["6 · Sat", "2026-10-24", "Spotify / website audio", "~55 s TTS voiceover", "4 · Sustainable Performance",
     "The One-Minute Neural Reset",
     "All personas",
     "Weekends are for recovery. Lets people *hear* the brand voice and links calm to the NEURALIS name.",
     "Completion rate", "assets/audio/neural-reset-voiceover.mp3"],
    ["7 · Sun", "2026-10-25", "Email newsletter", "Weekly digest", "1 + 2 + 4",
     "Three rules for a clearer Monday",
     "Selin, existing subscribers",
     "Sunday evening planning moment. Recaps the week's content and sets up Monday's habit loop.",
     "Open rate · click rate", "Links to Day 1–6 content"],
]


def build() -> None:
    docs = ROOT / "docs"
    docs.mkdir(exist_ok=True)

    with open(docs / "content-calendar.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(ROWS)

    wb = Workbook()
    ws = wb.active
    ws.title = "7-Day Calendar"
    ws.append(HEADERS)
    for r in ROWS:
        ws.append(r)

    head_fill = PatternFill("solid", fgColor="0B0F14")
    alt_fill = PatternFill("solid", fgColor="EEF9FB")
    thin = Side(style="thin", color="CBD5E1")
    for c in ws[1]:
        c.font = Font(bold=True, color="00F2FE")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if i % 2 == 0:
                c.fill = alt_fill
    widths = [9, 12, 20, 20, 22, 46, 22, 52, 28, 34]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # second sheet: format coverage vs. brief requirements
    cov = wb.create_sheet("Format Coverage")
    cov.append(["Required format (brief, Stage 03)", "Delivered", "Calendar day", "File"])
    for r in [
        ["SEO blog post", "1", "Day 3", "content/blog-post.md"],
        ["Social posts written with RTF", "3", "Days 1, 2, 4", "content/social-posts.md"],
        ["On-palette brand images", "3", "Days 2, 3, 5", "assets/images/"],
        ["15–20 s promo video", "1 (18 s)", "Day 5", "assets/video/neuralis-teaser-9x16.mp4"],
        ["On-brand TTS voiceover", "1", "Day 6", "assets/audio/neural-reset-voiceover.mp3"],
    ]:
        cov.append(r)
    for c in cov[1]:
        c.font = Font(bold=True, color="00F2FE")
        c.fill = head_fill
    for i, wdt in enumerate([36, 12, 16, 44], start=1):
        cov.column_dimensions[get_column_letter(i)].width = wdt

    wb.save(docs / "content-calendar.xlsx")
    print("Wrote docs/content-calendar.xlsx and docs/content-calendar.csv")


if __name__ == "__main__":
    build()
